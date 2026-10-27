"""엑셀(xlsx) -> Piece 목록. 시트마다 머리 행 아래의 행을 GROUP_ROWS 행 이하로 고르게 묶어 Piece 하나로 만든다. SCRUM-109.
(GROUP_ROWS 가 10 이면 11행은 10+1 이 아니라 6+5). 지금은 1, 행 하나가 Piece 하나다:
10행 묶음과 비교해 표 질문이 크게 좋아졌다 (SCRUM-110, 운영계획서 엑셀 질문 22개, Hit@1 59% → 77%)

  머리 행     시트에서 처음으로 글이 있는 행. 각 열의 이름으로 쓴다
  행 한 줄    "열 이름: 값 | 열 이름: 값" (워드 표와 같은 모양). 빈 칸은 뺀다
  location   {"type": "sheet", "sheet": 시트 이름, "rows": [시작 행, 끝 행], "label": "1분기 시트 12행"}  (RAG 명세)
             행 번호는 엑셀 화면에 보이는 번호 그대로다
  context    시트 이름. 검색할 때 조각 앞에 붙는다

  병합된 칸은 병합 범위의 모든 칸에 같은 값을 채운다 (행마다 뜻이 통하게)
  숫자는 엑셀에 보이는 모양에 가깝게: 천 단위 쉼표, 백분율, 날짜
  수식 칸은 엑셀이 저장해 둔 계산 결과를 쓴다. 결과가 저장돼 있지 않으면(엑셀로 한 번도 열지 않은 파일) 그 칸은 빠진다
  숨긴 시트는 읽지 않는다 (계산용·메모가 근거로 나오지 않게)
읽지 못하는 것: 차트, 그림, 칸 메모. 한 시트에 표가 여러 개이거나 머리 행이 두 줄이면 첫 행만 머리 행으로 본다.
글이 없으면 no_text, 열 수 없으면 unsupported_format, 칸이 상한을 넘으면 too_large. 꺼낸 글은 로그·오류 글에 남기지 않는다.
"""
import io
import re
from datetime import date, datetime, time

from ..errors import NO_TEXT, TOO_LARGE, UNSUPPORTED_FORMAT, IngestError
from ..models import Piece
from .minutes import LoadResult
from .text import normalize

GROUP_ROWS = 1                   # 조각 하나에 넣는 최대 행 수. 10·5·1 을 비교해 정함 (SCRUM-110)
MAX_CELLS = 200_000              # 정해 둔 값. SCRUM-110 에서 재서 조정한다
_DECIMALS = re.compile(r"0\.(0+)")


def load_xlsx(data: bytes) -> LoadResult:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), data_only=True)
    except Exception:                       # 깨진 zip, 빠진 부분, XML 오류 등. 내용은 오류에 남기지 않는다
        raise IngestError(UNSUPPORTED_FORMAT, "열 수 없는 엑셀 파일") from None
    try:
        cells = sum(ws.max_row * ws.max_column for ws in wb.worksheets)
        if cells > MAX_CELLS:
            raise IngestError(TOO_LARGE, f"칸 {cells}개, 상한 {MAX_CELLS}개")
        pieces, skipped = [], 0
        for ws in wb.worksheets:
            if ws.sheet_state != "visible":
                skipped += 1
                continue
            sheet_pieces = _sheet_pieces(ws)
            skipped += not sheet_pieces
            pieces += sheet_pieces
    finally:
        wb.close()
    if not pieces:
        raise IngestError(NO_TEXT, "보이는 시트에 글이 없음")
    return LoadResult(pieces, skipped)


def _sheet_pieces(ws) -> list[Piece]:
    merged = {}
    for rng in ws.merged_cells.ranges:      # 병합 범위의 모든 칸에 왼쪽 위 칸의 값을 채운다
        top = ws.cell(rng.min_row, rng.min_col)
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                merged[(r, c)] = top

    rows = []                               # (행 번호, [칸 글])
    for row in ws.iter_rows():
        values = [cell_text(merged.get((cell.row, cell.column), cell)) for cell in row]
        if any(values):
            rows.append((row[0].row, values))
    if not rows:
        return []
    (header_row, header), data = rows[0], rows[1:]
    if not data:
        return [_piece(ws.title, header_row, header_row, [" | ".join(v for v in header if v)])]
    pieces = []
    for group in even_groups(data, GROUP_ROWS):
        lines = [" | ".join(f"{header[j]}: {v}" if j < len(header) and header[j] else v
                            for j, v in enumerate(values) if v) for _, values in group]
        pieces.append(_piece(ws.title, group[0][0], group[-1][0], lines))
    return pieces


def even_groups(items, most):
    """most 개 이하로, 크기가 고르게 나눈다. 11개를 10 이하로 → 6, 5"""
    count = -(-len(items) // most)
    size, extra = divmod(len(items), count)
    out, start = [], 0
    for i in range(count):
        end = start + size + (i < extra)
        out.append(items[start:end])
        start = end
    return out


def _piece(sheet, first, last, lines) -> Piece:
    rows = f"{first}행" if first == last else f"{first}~{last}행"
    return Piece(content=normalize("\n".join(lines)), kind="sheet", context=sheet,
                 location={"type": "sheet", "sheet": sheet, "rows": [first, last], "label": f"{sheet} 시트 {rows}"})


def cell_text(cell) -> str:
    """칸 값을 엑셀에 보이는 모양에 가깝게. 빈 칸·계산 결과가 없는 수식은 ""."""
    v = cell.value
    if v is None:
        return ""
    fmt = cell.number_format or "General"
    if isinstance(v, datetime):
        return v.date().isoformat() if v.time() == time(0) else v.strftime("%Y-%m-%d %H:%M")
    if isinstance(v, (date, time)):
        return v.isoformat()
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, (int, float)):
        return _number(v, fmt)
    return normalize(str(v)).replace("\n", " ")


def _number(v, fmt) -> str:
    m = _DECIMALS.search(fmt)
    decimals = len(m.group(1)) if m else 0
    if "%" in fmt:
        return f"{v * 100:.{decimals}f}%"
    if "," in fmt:
        return f"{v:,.{decimals}f}"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)
